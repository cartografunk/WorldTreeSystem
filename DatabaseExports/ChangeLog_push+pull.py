from core.db import get_engine
from core.libs import pd
import openpyxl
import os

EXCEL_PATH = "masterdatabase_export.xlsx"
CHANGELOG_SHEET = "ChangeLog"

# Paso 1: Cargar ChangeLog si existe
if os.path.exists(EXCEL_PATH):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    if CHANGELOG_SHEET in wb.sheetnames:
        ws = wb[CHANGELOG_SHEET]
        data = list(ws.values)
        columns = data[0]
        changelog_df = pd.DataFrame(data[1:], columns=columns)
    else:
        changelog_df = pd.DataFrame(columns=[
            "contract_code", "target_field", "new_value", "requested_by", "timestamp", "note"
        ])
else:
    changelog_df = pd.DataFrame(columns=[
        "contract_code", "target_field", "new_value", "requested_by", "timestamp", "note"
    ])

# Paso 2: Exportar las tablas de la base como de costumbre
engine = get_engine()
query = """
SELECT table_name
FROM information_schema.tables
WHERE table_schema = 'masterdatabase'
ORDER BY table_name;
"""
tables = pd.read_sql(query, engine)["table_name"].tolist()

def remove_tz(df):
    for col in df.select_dtypes(include=["datetimetz"]).columns:
        df[col] = df[col].dt.tz_localize(None)
    return df

with pd.ExcelWriter(EXCEL_PATH, engine="openpyxl", mode="w") as writer:
    # Primero ChangeLog
    changelog_df.to_excel(writer, sheet_name=CHANGELOG_SHEET, index=False)
    # Luego todas las demás tablas
    for table in tables:
        print(f"Exportando {table} ...")
        df = pd.read_sql(f'SELECT * FROM masterdatabase."{table}"', engine)
        if not df.empty:
            df = remove_tz(df)
            df = df.sort_values(by=df.columns[0])
        df.to_excel(writer, sheet_name=table[:31], index=False)
    # Si quieres agregar FieldsCatalog aquí (opcional)

print("✅ Exportación completa. ChangeLog preservado.")
